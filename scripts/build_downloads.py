"""Build polished .xlsx downloads for the GitHub Pages dashboard."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "dashboard.json"
OUT_DIR = ROOT / "downloads"

COLUMNS = [
    ("period_label", "기간"),
    ("revenue", "매출액"),
    ("operating_income", "영업이익"),
    ("net_income", "당기순이익"),
    ("operating_cash_flow", "영업CF"),
    ("investing_cash_flow", "투자CF"),
    ("financing_cash_flow", "재무CF"),
    ("cash_and_cash_equivalents", "현금"),
    ("total_assets", "총자산"),
    ("total_liabilities", "총부채"),
    ("total_equity", "자본"),
    ("operating_margin_pct", "영업이익률"),
    ("net_margin_pct", "순이익률"),
    ("roa_pct", "ROA"),
    ("roe_pct", "ROE"),
    ("current_ratio_pct", "유동비율"),
    ("debt_to_equity_pct", "부채/자본"),
    ("dso_days", "DSO"),
]

MONEY_FIELDS = {
    "revenue", "operating_income", "net_income", "operating_cash_flow",
    "investing_cash_flow", "financing_cash_flow", "cash_and_cash_equivalents",
    "total_assets", "total_liabilities", "total_equity",
}
PCT_FIELDS = {
    "operating_margin_pct", "net_margin_pct", "roa_pct", "roe_pct",
    "current_ratio_pct", "debt_to_equity_pct",
}

def cell_value(field: str, value):
    if value is None:
        return None
    if field in MONEY_FIELDS:
        return value / 100_000_000
    return value

def build(rows: list[dict], title: str, filename: str, updated: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title.replace("-", " ")

    ws["A1"] = f"SD BIOSENSOR — {title}"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="16324A")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    ws["A2"] = "Source: OpenDART · 금액 단위: 억원 · 비율 단위: % · DSO: days"
    ws["A2"].font = Font(color="627D98", italic=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
    ws["A3"] = f"Dashboard updated_at: {updated}"
    ws["A3"].font = Font(color="627D98", size=9)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(COLUMNS))

    header_row = 5
    for col, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(header_row, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0969AD")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, row in enumerate(rows, start=header_row + 1):
        for col_idx, (field, _) in enumerate(COLUMNS, start=1):
            value = row.get(field)
            if field == "period_label" and not value:
                value = str(row.get("year", ""))
            value = cell_value(field, value)
            cell = ws.cell(row_idx, col_idx, value)
            cell.alignment = Alignment(horizontal="left" if field == "period_label" else "right")
            if field in MONEY_FIELDS:
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
            elif field in PCT_FIELDS:
                cell.number_format = '0.00;[Red](0.00);-'
            elif field == "dso_days":
                cell.number_format = '0.0;[Red](0.0);-'

    last_row = header_row + len(rows)
    last_col = len(COLUMNS)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = "B6"
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[5].height = 22

    widths = [14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 12, 12, 10, 10, 12, 12, 10]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_DIR / filename)

with DATA_PATH.open(encoding="utf-8") as f:
    data = json.load(f)

for key, title, filename in [
    ("annual", "Annual", "SD_Biosensor_annual.xlsx"),
    ("half_year", "Half-year", "SD_Biosensor_half_year.xlsx"),
    ("quarterly", "Quarterly", "SD_Biosensor_quarterly.xlsx"),
]:
    build(data["tables"].get(key, []), title, filename, data.get("updated_at", ""))

print("Excel 다운로드 파일 생성 완료.")
